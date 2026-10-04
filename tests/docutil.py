"""Builds small PDF and DOCX files for tests."""

import io


def make_pdf(pages: list[str], title: str | None = None) -> bytes:
    """A minimal valid PDF with one text line per page (Helvetica, WinAnsi)."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    pages_id = len(objects) + 1 + 2 * len(pages) + (1 if title else 0) + 1
    page_ids = []
    for text in pages:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("cp1252")
        stream = b"BT /F1 12 Tf 72 720 Td (" + escaped + b") Tj ET"
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
                b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (pages_id, content, font)
            )
        )
    info = add(b"<< /Title (" + title.encode("cp1252") + b") >>") if title else None
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)
    kids = b" ".join(b"%d 0 R" % i for i in page_ids)
    assert add(b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(page_ids)) == pages_id
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    trailer = b"<< /Size %d /Root %d 0 R" % (len(objects) + 1, catalog)
    if info:
        trailer += b" /Info %d 0 R" % info
    out.write(b"trailer\n" + trailer + b" >>\nstartxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()


def make_docx(title: str, sections: list[tuple[str, str]]) -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.title = title
    for heading, text in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Key"
    table.rows[0].cells[1].text = "Wert"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
